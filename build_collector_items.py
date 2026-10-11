# -*- coding: utf-8 -*-
"""Build collector_items.json with precise find hints (place + person + method)."""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

ROOT = Path(r"D:\Unpacking game")
CONTENT = ROOT / r"afb02f5bbd8d0d94fe47e171a61cd51ee2ffeb73\Output\Exports\Octopath_Traveler\Content"
STORE_XLSX = Path(r"D:\Users\10129\Downloads\Octopath Traveler Resource.xlsx")

TOWN_CN = {
    "Flamesgrace": "弗雷姆葛雷斯",
    "Atlasdam": "阿特拉斯达姆",
    "Rippletide": "利布尔泰德",
    "Cobbleston": "鹅卵石村",
    "Sunshade": "桑谢德",
    "Clearbrook": "库利亚布鲁克",
    "Bolderfall": "柏达弗尔",
    "S'warkii": "希・沃尔基",
    "Stillsnow": "史迪尔斯诺",
    "Noblecourt": "诺布尔寇德",
    "Goldshore": "葛鲁德修亚",
    "Stonegard": "斯托冈德",
    "Wellspring": "威尔斯宾克",
    "Saintsbridge": "圣特布里吉",
    "Quarrycrest": "库欧利克雷斯特",
    "Victors Hollow": "维克塔霍洛",
    "Northreach": "诺斯利奇",
    "Wispermill": "维斯帕米尔",
    "Grandport": "格兰波特",
    "Everhold": "艾巴霍多",
    "Marsalim": "马尔沙利姆",
    "Riverford": "利维福德",
    "Orewell": "欧亚威尔",
    "Duskbarrow": "达斯克瓦罗",
}

INFO_REG = {
    "SNO": "霜地", "PLA": "平原", "SEA": "海岸", "MOU": "高原",
    "DES": "沙漠", "RIV": "河流", "CLI": "崖地", "FOR": "森林",
}


def load_rows(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))[0]["Rows"]


def text_value(text_rows: dict, key: str) -> str:
    if not key or key == "None":
        return ""
    row = text_rows.get(key)
    if not row:
        return ""
    t = row.get("Text") or {}
    return t.get("LocalizedString") or t.get("SourceString") or ""


def field(row, prefix: str):
    if not isinstance(row, dict):
        return None
    for k, v in row.items():
        if k.startswith(prefix):
            return v
    return None


def map_candidates(region: str, area_code: str) -> list[str]:
    out = []
    if not region or not area_code:
        return out
    kind, rest = area_code[0], area_code[1:]
    m = re.match(r"^(\d)(\d?)([A-Za-z]*)$", rest)
    if kind in ("D", "F") and m:
        level, sub, suffix = m.group(1), m.group(2), m.group(3)
        if suffix.lower() == "job":
            out.append(f"{kind}_{region}_{level}_Job")
        base = f"{kind}_{region}_{level}_{sub}" if sub else f"{kind}_{region}_{level}"
        if suffix and suffix.lower() != "job":
            su = suffix.upper()
            out.append(f"{base}_{su}")
            out.append(f"{base}{su}")
        out.append(base)
    if kind == "T":
        if len(rest) == 1:
            out.append(f"T_{region}_{rest}")
        elif len(rest) == 2:
            out.append(f"T_{region}_{rest[0]}_{rest[1]}")
            out.append(f"T_{region}_{rest}")
        elif len(rest) == 3:
            out.append(f"T_{region}_{rest[0]}_{rest[1]}_{rest[2]}")
            out.append(f"T_{region}_{rest}")
    return out


def object_map_name(text_rows: dict, object_id: str) -> str:
    parts = object_id.split("_")
    if len(parts) < 3:
        return ""
    region, area = parts[1], parts[2]
    for cand in map_candidates(region, area):
        name = text_value(text_rows, cand)
        if name:
            return name
    return f"{region}_{area}"


def info_place_hint(item_key: str) -> str:
    m = re.match(r"^INFO_([A-Z]{3})_([A-Z0-9]+?)(?:_(\d+))?$", item_key)
    if not m:
        return ""
    reg = INFO_REG.get(m.group(1), m.group(1))
    mid = m.group(2)
    if mid.startswith("T"):
        place = "城镇"
        if "S" in mid:
            place += "（小）"
        elif "M" in mid:
            place += "（中）"
        elif "L" in mid:
            place += "（大）"
    elif mid.startswith("F"):
        place = "野外道路"
    elif mid.startswith("D"):
        place = "迷宫"
    else:
        place = mid
    return f"打听/探查：{reg}{place}"


def norm_item_name(s: str) -> str:
    s = (s or "").strip()
    s = re.sub(r"^[\*＊]+\s*", "", s)
    s = re.sub(r"^\(\*\)\s*", "", s)
    s = s.replace("’", "'").replace("‘", "'")
    s = s.replace("'", "").replace(".", "")
    s = s.casefold()
    s = s.replace("nourshing", "nourishing")
    s = s.replace("myterious", "mysterious")
    s = s.replace("aristrocrat", "aristocrat")
    s = re.sub(r"\s*\((m|l)\)", r" \1", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def display_item_name(s: str) -> str:
    s = (s or "").strip()
    return re.sub(r"\s*\(\d+(?:=\d+)+\)$", "", s).strip()


def npc_base_name(s: str) -> str:
    s = (s or "").strip()
    s = re.sub(r"\s*\(\d+\)\s*$", "", s)
    return s


def load_lang_en_to_cn() -> dict[str, str]:
    out = dict(TOWN_CN)
    if not STORE_XLSX.exists():
        return out
    import openpyxl
    wb = openpyxl.load_workbook(STORE_XLSX, read_only=True, data_only=True)
    ws = wb["Language"]
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 2 or not row[1]:
            continue
        en = str(row[1]).strip()
        cn = None
        if len(row) > 7 and row[7]:
            cn = str(row[7]).strip()
        elif len(row) > 6 and row[6]:
            cn = str(row[6]).strip()
        if en and cn:
            out[en] = cn
    return out


# Excel Language 页常用「警卫」，游戏内镇民几乎都叫「卫兵」。
NPC_EN_OFFICIAL_CN = {
    "Guard": "卫兵",
}


def loc_cn(en: str, lang: dict[str, str]) -> str:
    en = (en or "").strip()
    if not en:
        return ""
    head = en.split("(")[0].strip()
    if head in NPC_EN_OFFICIAL_CN:
        return NPC_EN_OFFICIAL_CN[head]
    cn = lang.get(en) or lang.get(head) or TOWN_CN.get(en) or TOWN_CN.get(head) or en
    cn = str(cn).replace("／", "/").replace("\r", "")
    cn = cn.split("\n")[0].split("/")[0].strip()
    return cn


def split_cell_items(raw) -> list[str]:
    if raw is None:
        return []
    s = str(raw).strip()
    if not s or s.casefold() in ("none", "x", "false", "true"):
        return []
    parts = re.split(r"[\n/]+", s)
    out = []
    for p in parts:
        p = p.strip()
        if not p or p.casefold() in ("none", "a hidden item", "hidden item"):
            continue
        if p.endswith("$") or re.match(r"^[\d\s,]+\$", p):
            continue
        out.append(p)
    return out


def build_indoor_lookup(text_rows: dict) -> set[tuple[str, str]]:
    """(NPC中文名, 地图中文名) that only/mostly spawn indoors (HouseInFlag)."""
    place_list = load_rows(CONTENT / "Character/Database/PlacementInfoList.json")
    place_data = load_rows(CONTENT / "Character/Database/PlacementInfoData.json")
    pid_to_map: dict[str, str] = {}
    for mk, row in place_list.items():
        labels = field(row, "LabelList_") or []
        mname = text_value(text_rows, mk) or mk
        for lab in labels:
            if lab:
                pid_to_map[lab] = mname
    counts: dict[tuple[str, str], list[int]] = defaultdict(lambda: [0, 0])
    for pid, row in place_data.items():
        name_id = row.get("TextLabel") if isinstance(row, dict) else None
        npc = text_value(text_rows, name_id) if name_id else ""
        if not npc:
            continue
        mname = pid_to_map.get(pid, "")
        if not mname:
            continue
        indoor = bool(row.get("HouseInFlag")) if isinstance(row, dict) else False
        counts[(npc, mname)][0 if indoor else 1] += 1
    out: set[tuple[str, str]] = set()
    for key, (inn, outd) in counts.items():
        if inn and inn >= outd:
            out.add(key)
    return out


def format_who(npc: str, place: str, indoor: set[tuple[str, str]]) -> str:
    if not npc:
        return place
    tagged = False
    if (npc, place) in indoor:
        tagged = True
    else:
        head = place.split(" -")[0].split("（")[0].strip()
        if head and (npc, head) in indoor:
            tagged = True
    if tagged and "房中" not in place and "酒馆" not in place:
        return f"{npc}（{place}，房中）"
    return f"{npc}（{place}）"


def parse_npc_sheet(lang: dict[str, str], indoor: set[tuple[str, str]]) -> dict[str, list[str]]:
    """norm EN item name -> hint lines with NPC + place."""
    out: dict[str, list[str]] = defaultdict(list)
    if not STORE_XLSX.exists():
        return out
    import openpyxl
    wb = openpyxl.load_workbook(STORE_XLSX, read_only=True, data_only=True)
    ws = wb["NPC"]
    place = ""
    npc = ""
    skip_headers = {"NPC", "Tier 1", "Tier 2", "Tier 3"}
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        b = cells[1] if len(cells) > 1 else None
        d = cells[3] if len(cells) > 3 else None
        f = cells[5] if len(cells) > 5 else None
        drop = cells[14] if len(cells) > 14 else None
        steal_items = split_cell_items(d)
        info_items = split_cell_items(f)
        drop_items = split_cell_items(drop)
        # Town/map headers have blank item columns. NPC rows with "X"/"None"
        # must not overwrite the current place (e.g. Kevin in Marsalim Palace).
        item_cells_filled = any(
            c is not None and str(c).strip()
            for c in (d, f, drop)
        )

        if isinstance(b, str) and b.strip() and not item_cells_filled:
            raw = b.strip()
            if raw in skip_headers:
                continue
            place = loc_cn(raw, lang)
            npc = ""
            continue

        if isinstance(b, str) and b.strip():
            npc = loc_cn(npc_base_name(b.strip()), lang)

        if not npc or not place:
            continue
        who = format_who(npc, place, indoor)
        for it in steal_items:
            k = norm_item_name(it)
            line = f"购买/偷取：{who}"
            if line not in out[k]:
                out[k].append(line)
        for it in drop_items:
            k = norm_item_name(it)
            line = f"比试/唆使掉落：{who}"
            if line not in out[k]:
                out[k].append(line)
        for it in info_items:
            k = norm_item_name(it)
            line = f"打听/探查：{who}"
            if line not in out[k]:
                out[k].append(line)
    return out


def parse_quest_rewards(lang: dict[str, str]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = defaultdict(list)
    if not STORE_XLSX.exists():
        return out
    import openpyxl
    wb = openpyxl.load_workbook(STORE_XLSX, read_only=True, data_only=True)
    ws = wb["Quests"]
    place = ""
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        b = cells[1] if len(cells) > 1 else None
        name = cells[2] if len(cells) > 2 else None
        npc = cells[3] if len(cells) > 3 else None
        reward = cells[5] if len(cells) > 5 else None
        if isinstance(b, str) and b.strip() and not name:
            raw = b.strip()
            if raw in ("Outside Starting Town",):
                place = "初始城镇附近"
            else:
                place = loc_cn(raw, lang)
            continue
        if not name:
            continue
        qn = loc_cn(str(name).strip(), lang)
        who = loc_cn(npc_base_name(str(npc).strip()), lang) if npc else ""
        locbit = f"（{place}）" if place else ""
        npbit = f"，NPC：{who}" if who else ""
        for it in split_cell_items(reward):
            k = norm_item_name(it)
            line = f"支线奖励：{qn}{locbit}{npbit}"
            if line not in out[k]:
                out[k].append(line)
    return out


def load_store_by_en_name() -> dict[str, list[str]]:
    out: dict[str, list[str]] = defaultdict(list)
    if not STORE_XLSX.exists():
        return out
    import openpyxl
    wb = openpyxl.load_workbook(STORE_XLSX, read_only=True, data_only=True)
    ws = wb["Store"]
    towns = []
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        header_towns = []
        for i in range(1, len(cells), 2):
            v = cells[i]
            if isinstance(v, str) and v.strip() in TOWN_CN:
                header_towns.append((i, TOWN_CN[v.strip()]))
        if len(header_towns) >= 4:
            towns = header_towns
            continue
        if not towns:
            continue
        for col, town in towns:
            if col >= len(cells) or not isinstance(cells[col], str):
                continue
            raw = cells[col]
            starred = raw.strip().startswith("(*)")
            key = norm_item_name(raw)
            if not key:
                continue
            label = f"{town}" + ("（章节后补货）" if starred else "")
            if label not in out[key]:
                out[key].append(label)
    return out


def shop_hint_line(towns: list[str]) -> str:
    if not towns:
        return ""
    extra = any("章节后补货" in t for t in towns)
    clean = [t.replace("（章节后补货）", "") for t in towns]
    seen = []
    for t in clean:
        if t not in seen:
            seen.append(t)
    if len(seen) >= 8:
        line = f"城镇商店可购：{'、'.join(seen[:6])}等（共{len(seen)}处）"
    else:
        line = "城镇商店：" + "、".join(seen)
    if extra:
        line += "（部分为章节后补货）"
    return line


def build_enemy_spawns(text_rows: dict) -> dict[str, list[str]]:
    """enemy_key -> unique CN map names from LevelTable volumes."""
    level_rows = load_rows(CONTENT / "Level/Database/LevelTable.json")
    enc_vol = load_rows(CONTENT / "Battle/Database/EncountVolumeData.json")
    encounter = load_rows(CONTENT / "Battle/Database/BattleEncounterData.json")
    groups = load_rows(CONTENT / "Battle/Database/EnemyGroupData.json")

    evm_to_enc: dict[str, set[str]] = defaultdict(set)
    for evm, row in enc_vol.items():
        lst = field(row, "EncounterList_") or []
        for item in lst:
            enc = field(item, "EncounterName_") if isinstance(item, dict) else None
            if enc and enc != "None":
                evm_to_enc[evm].add(enc)

    enc_to_ene: dict[str, set[str]] = defaultdict(set)
    for enc, row in encounter.items():
        glst = field(row, "Group_") or []
        for g in glst:
            gid = field(g, "GroupID_") if isinstance(g, dict) else None
            if not gid or gid == "None" or gid not in groups:
                continue
            for ene in field(groups[gid], "EnemyID_") or []:
                if ene and ene != "None":
                    enc_to_ene[enc].add(ene)

    spawns: dict[str, list[str]] = defaultdict(list)
    seen: dict[str, set[str]] = defaultdict(set)
    for row in level_rows.values():
        mk = row.get("LevelName")
        if not mk or mk == "None":
            continue
        mname = text_value(text_rows, mk) or mk
        for evm in row.get("EncountVolumeNameList") or []:
            if not evm or evm == "None":
                continue
            for enc in evm_to_enc.get(evm, ()):
                for ene in enc_to_ene.get(enc, ()):
                    if mname not in seen[ene]:
                        seen[ene].add(mname)
                        spawns[ene].append(mname)
    return spawns


def uniq_join(parts: list[str], limit: int = 12) -> str:
    seen = []
    for p in parts:
        p = (p or "").strip()
        if p and p not in seen:
            seen.append(p)
        if len(seen) >= limit:
            break
    return "；".join(seen)


def pick_lines(lines: list[str], limit: int) -> list[str]:
    seen = []
    for x in lines:
        if x not in seen:
            seen.append(x)
        if len(seen) >= limit:
            break
    return seen


def main() -> None:
    inv = json.loads((ROOT / "collector_item_inventory.json").read_text(encoding="utf-8"))["items"]
    special = json.loads((ROOT / "collector_special_hints.json").read_text(encoding="utf-8"))
    item_rows = load_rows(CONTENT / "Item/Database/ItemDB.json")
    text_rows = load_rows(CONTENT / "GameText/Database/GameTextZH_CN.json")
    object_rows = load_rows(CONTENT / "Object/Database/ObjectData.json")
    enemy_rows = load_rows(CONTENT / "Character/Database/EnemyDB.json")
    hear_rows = load_rows(CONTENT / "Character/Database/NPCHearInfoData.json")

    lang = load_lang_en_to_cn()
    indoor = build_indoor_lookup(text_rows)
    npc_by_item = parse_npc_sheet(lang, indoor)
    print(f"indoor_npc_places={len(indoor)}")
    quest_by_item = parse_quest_rewards(lang)
    store_by_en = load_store_by_en_name()
    spawns = build_enemy_spawns(text_rows)
    print(f"lang={len(lang)} npc_items={len(npc_by_item)} quest_items={len(quest_by_item)} store={len(store_by_en)} spawn_enemies={len(spawns)}")

    id_to_key = {}
    for k, v in item_rows.items():
        iid = field(v, "ItemID_")
        if iid is not None:
            id_to_key[int(iid)] = k

    chests: dict[str, list[str]] = defaultdict(list)
    hidden: dict[str, list[str]] = defaultdict(list)
    for oid, row in object_rows.items():
        label = row.get("HaveItemLabel") or "None"
        if label in ("None", None) or row.get("IsMoney"):
            continue
        loc = object_map_name(text_rows, oid)
        if oid.startswith("TrsrBx_"):
            line = f"宝箱：{loc}" if loc else "宝箱"
            if line not in chests[label]:
                chests[label].append(line)
        elif oid.startswith("HdnPnt_"):
            line = f"隐藏道具：{loc}" if loc else "隐藏道具"
            if line not in hidden[label]:
                hidden[label].append(line)

    drops: dict[str, list[str]] = defaultdict(list)
    for ek, ev in enemy_rows.items():
        have = field(ev, "HaveItemID_")
        if not have or have == "None":
            continue
        name_id = field(ev, "DisplayNameID_")
        ename = text_value(text_rows, name_id) or ek
        maps = spawns.get(ek, [])
        if maps:
            shown = "、".join(maps[:3])
            extra = f"等{len(maps)}处" if len(maps) > 3 else ""
            tag = f"野怪偷取/掉落：{ename}（{shown}{extra}）"
        else:
            # Town/NPC battles have no field volume — skip nameless location.
            continue
        if tag not in drops[have]:
            drops[have].append(tag)

    hear: dict[str, int] = defaultdict(int)
    for _hk, hv in hear_rows.items():
        iid = field(hv, "ItemID_")
        if iid and iid != "None":
            hear[iid] += 1

    vague = 0
    out = []
    for rec in inv:
        iid = int(rec["id"])
        if iid <= 0:
            continue
        key = id_to_key.get(iid, "")
        en = rec.get("en") or ""
        nk = norm_item_name(en)
        keys = {nk}
        cn_name = rec.get("name") or ""
        for en_name, cn in lang.items():
            if cn == cn_name:
                keys.add(norm_item_name(en_name))
        npc_lines = []
        quest_lines = []
        shop_towns = []
        for k in keys:
            npc_lines.extend(npc_by_item.get(k, []))
            quest_lines.extend(quest_by_item.get(k, []))
            shop_towns.extend(store_by_en.get(k, []))
        bits = []
        sp = special.get(str(iid), {})
        if sp.get("hint"):
            bits.append(sp["hint"])
        omit = list(sp.get("omit_npc_substrings") or sp.get("drop_npc_substrings") or [])
        if omit:
            npc_lines = [ln for ln in npc_lines if not any(s in ln for s in omit)]

        bits.extend(pick_lines(npc_lines, 8))
        bits.extend(pick_lines(quest_lines, 4))
        shop = shop_hint_line(shop_towns)
        if shop:
            bits.append(shop)

        if key:
            bits.extend(pick_lines(chests.get(key, []), 5))
            bits.extend(pick_lines(hidden.get(key, []), 4))
            bits.extend(pick_lines(drops.get(key, []), 4))
            if key.startswith("INFO_"):
                bits.append(info_place_hint(key))
                n = hear.get(key, 0)
                if n and not npc_lines:
                    bits.append(f"可由 {n} 处打听/探查取得")

        hint = uniq_join(bits, 12)
        if not hint:
            vague += 1
            if rec.get("type") == 14:
                hint = "知识类：请在对应城镇/迷宫打听或探查（全资源清单 NPC 页 Scrutinize/Inquire 列）。"
            elif rec.get("type") == 13:
                hint = "贵重物品：多为剧情/支线关键道具，对照全资源清单 Quests 与 NPC 购买/偷取列。"
            else:
                hint = "未在商店表、NPC 购买/偷取/比试掉落、宝箱或隐藏点中匹配到来源，请用英文名对照全资源清单 NPC/Store/Quests。"

        out.append({
            "kind": "Collector",
            "item_id": iid,
            "item_key": key,
            "name": display_item_name(rec["name"]),
            "type": rec["type"],
            "type_name": rec["type_name"],
            "find_hint": hint,
            "missable_note": sp.get("missable_note") or "",
        })

    dest = ROOT / "collector_items.json"
    dest.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"wrote {len(out)} items, unmatched={vague} -> {dest}")


if __name__ == "__main__":
    main()
